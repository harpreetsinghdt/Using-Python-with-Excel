from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Reference, Series, PieChart3D

wb = load_workbook('../crime_report.xlsx')
ws = wb.active

chart = BarChart()
data = Reference(ws, min_row=10,min_col=1, max_row=13,  max_col=13)
labels = Reference(ws, min_col=2, min_row=8, max_row=8, max_col=13)
chart.add_data(data, from_rows=True ,titles_from_data=True)
chart.set_categories(labels)
chart.title = 'Counterfeit Crimes by District'
chart.height = 4.56
chart.width = 20.3
ws.add_chart(chart, 'B14')

chart2 = PieChart()
labels = Reference(ws, min_col=15, min_row=7, max_col=16, max_row=7)
data = Reference(ws, min_col=15, min_row=8, max_row=8,max_col=16)
chart2.add_data(data, from_rows=True)
chart2.set_categories(labels)
chart2.title = '% Counterfeit Crimes'
chart2.height = 4.56
chart2.width = 8.45
ws.add_chart(chart2, 'N14')

wb.save('lines.xlsx')

